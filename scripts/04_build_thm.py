"""04_build_thm.py - builds the Tourism and Hospitality Management (THM) submission file:
pandoc renders paper/paper_thm_body.md with the THM reference styles, then python-docx adds the
title, the author/abstract block and the author biographies required by the journal template."""
import subprocess, os, copy
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__)); P = os.path.join(HERE, "..", "paper")
os.chdir(P)
ANON = os.environ.get("ANON") == "1"   # ANON=1: anonymised manuscript for blind review
SRC = "paper_thm_body.md"
if ANON:
    import re
    txt = open(SRC, encoding="utf-8").read()
    # drop the acknowledgements section (names institutions and co-author roles)
    txt = re.sub(r"# ACKNOWLEDGEMENTS\n.*?(?=\n# )", "", txt, flags=re.S)
    # self-citation of the data/code archive
    txt = txt.replace("code and data are openly available (Uroš et al., 2026)",
                      "code and data are archived in a public repository (Authors, 2026; DOI provided at acceptance)")
    txt = txt.replace("(Uroš et al., 2026)", "(Authors, 2026)")
    txt = re.sub(r"^Uroš, V\., Ćaćić, M\., & Sorić, I\. \(2026\)\..*?\n\n", "", txt, flags=re.M | re.S)
    txt = re.sub(r"(# REFERENCES\n\n)", r"\1Authors. (2026). *Data set and code* [Details omitted for blind review].\n\n", txt)
    assert "Uroš" not in txt and "github.com" not in txt and "vedranu" not in txt, "identifying text remains"
    SRC = "_thm_body_anon.md"; open(SRC, "w", encoding="utf-8").write(txt)
subprocess.run(["pandoc", SRC, "-o", "_thm_body.docx", "--resource-path=.", "--columns=60", "--reference-doc=thm_reference.docx"], check=True)
if ANON: os.remove(SRC)

TITLE = 'FROM TRANSIT STOP TO DESTINATION? A DECADE OF E-VISITOR DATA ON TOURISM IN KNIN, CROATIA'
AUTHORS = [
    dict(name="Vedran UROŠ", corr=True, orcid="https://orcid.org/0009-0007-6568-0827",
         inst=['University of Applied Sciences "Marko Marulić" in Knin', 'Faculty of Organization and Informatics, University of Zagreb'], country="Croatia", email="vuros@veleknin.hr"),
    dict(name="Marijo ĆAĆIĆ", corr=False, orcid="https://orcid.org/0009-0008-7197-4263",
         inst=["Faculty of Food Technology and Biotechnology, University of Zagreb", "The Town of Knin Administration"], country="Croatia", email="mcacic@pbf.hr"),
    dict(name="Ivan SORIĆ", corr=False, orcid="https://orcid.org/0009-0009-1699-1973",
         inst=["Tourist Board of the Town of Knin", "University of Zadar (doctoral programme)"], country="Croatia", email="director@tz-knin.hr"),
]
ABSTRACT = [
    ("Purpose", "The paper examines how tourism in Knin, a small town in the Dalmatian hinterland, developed during the first decade of Croatia's e-Visitor system, whether growth changed its seasonality and source-market structure, and what the data imply for 2030."),
    ("Methodology/Design/Approach", "The complete monthly e-Visitor record is analysed with STL decomposition, Gini coefficients, Herfindahl-Hirschman indices, segmented-trend break estimation with bootstrap inference, and SARIMA, ETS and Holt-Winters models validated by hold-out and rolling-origin evaluation and projected to 2030."),
    ("Findings", "Arrivals grew from 1,404 (2016) to 6,178 (2025) and overnight stays from 3,456 to 12,558. Post-pandemic growth was steeper and specific to Knin (county arrivals stayed below 2019), but break dates have wide confidence intervals. Seasonality did not fall; the summer peak flattened into a longer summer. Foreign demand diversified across 12-17 effective markets, but average stay fell to two nights. Growth stalled in 2026, so projections to 2030 range from 5,800 to 12,200 arrivals."),
    ("Originality of the research", "First decade-long monthly analysis of a single Croatian inland destination, with open data and code."),
]
KEYWORDS = "tourism demand, seasonality, source markets, structural break, forecasting, Croatia"
BIOS = [
    'Vedran Uroš is a senior lecturer at the University of Applied Sciences "Marko Marulić" in Knin, where he serves as acting vice-dean for research and international cooperation, and a doctoral candidate at the Faculty of Organization and Informatics, University of Zagreb. His research covers energy-efficient wireless sensor networks and UAV-assisted data collection, as well as applied data analysis in tourism, agribusiness and regional development. He teaches courses in information security, digital transformation and local and regional development.',
    "Marijo Ćaćić is a senior lecturer at the Faculty of Food Technology and Biotechnology, University of Zagreb, where he teaches Management and Marketing of Food Products in the graduate programme. He holds a master's degree in economics and is a doctoral student in Business Economics at the Faculty of Economics and Business, University of Rijeka. He also serves as Mayor of the Town of Knin, which gives him first-hand insight into local development and destination management. His professional interests include management, food marketing, and local and regional economic development.",
    "Ivan Sorić is Director of the Tourist Board of the Town of Knin, where he is responsible for destination management, promotion and the tourist registration system of the town. He holds a degree in philosophy and religious sciences and is a doctoral student at the University of Zadar in the doctoral programme The Adriatic - a Link between Continents. His professional interests include destination management, cultural and heritage tourism and the development of tourism in the Dalmatian hinterland.",
]

doc = Document("_thm_body.docx")
body = doc.element.body
first = body[0]

def para_before(ref, text="", style=None, bold=False, italic=False, size=None, align=None, font="Times New Roman"):
    p = doc.add_paragraph(); ref.addprevious(p._p)
    if style: p.style = doc.styles[style]
    if text:
        r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.name = font
        if size: r.font.size = Pt(size)
    if align is not None: p.alignment = align
    p.paragraph_format.space_after = Pt(4); p.paragraph_format.space_before = Pt(0)
    return p

def set_cell_borders(cell, val="nil"):
    tcPr = cell._tc.get_or_add_tcPr(); b = OxmlElement("w:tcBorders")
    for e in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{e}"); el.set(qn("w:val"), val); b.append(el)
    tcPr.append(b)

# title
tp = para_before(first, TITLE, bold=True, size=14, align=WD_ALIGN_PARAGRAPH.LEFT, font="Arial"); tp.paragraph_format.space_after = Pt(12)
# author / abstract table
tbl = doc.add_table(rows=1, cols=2); first.addprevious(tbl._tbl)
tbl.autofit = False
left, right = tbl.rows[0].cells
left.width = Cm(5.2); right.width = Cm(12.4)
for c in (left, right): set_cell_borders(c)
left.paragraphs[0].text = ""
def add_p(cell, text, bold=False, italic=False, size=9, after=2, prefix_italic=None):
    p = cell.add_paragraph(); p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(0)
    if prefix_italic:
        r = p.add_run(prefix_italic); r.italic = True; r.font.size = Pt(size); r.font.name = "Times New Roman"
    r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size); r.font.name = "Times New Roman"
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if cell is right else WD_ALIGN_PARAGRAPH.LEFT
    return p
for a in ([] if ANON else AUTHORS):
    add_p(left, a["name"] + (", " if a["corr"] else ""), bold=True, size=10, after=0)
    if a["corr"]: add_p(left, "Corresponding Author", size=9, after=0)
    add_p(left, a["orcid"], size=9, after=0)
    for i in a["inst"]: add_p(left, i, size=9, after=0)
    add_p(left, a["country"], size=9, after=0)
    add_p(left, a["email"], size=9, after=8)
right.paragraphs[0].text = ""
add_p(right, "Abstract", bold=True, size=10, after=4)
for k, v in ABSTRACT: add_p(right, " - " + v, size=9, after=4, prefix_italic=k)
kp = add_p(right, KEYWORDS, size=9, after=0); r0 = kp.runs[0]; r0.text = " " + KEYWORDS; rk = copy.deepcopy(r0._r); kp._p.insert(list(kp._p).index(r0._r), rk); kp.runs[0].text = "Keywords"; kp.runs[0].bold = True
# empty paragraph after table
para_before(first, "")

# biographies at end: one-column bordered table
bt = doc.add_table(rows=0 if ANON else len(BIOS), cols=1); bt.style = doc.styles["Table Grid"] if "Table Grid" in [s.name for s in doc.styles] else bt.style
for i, b in enumerate([] if ANON else BIOS):
    c = bt.rows[i].cells[0]; c.paragraphs[0].text = ""
    p = c.paragraphs[0]; name, rest = b.split(" is ", 1)
    r = p.add_run(name); r.bold = True; r.font.size = Pt(9); r.font.name = "Times New Roman"
    r = p.add_run(" is " + rest); r.font.size = Pt(9); r.font.name = "Times New Roman"; p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
# keep table/figure captions with the following element
for p in doc.paragraphs:
    t = p.text.strip()
    if (t.startswith("Table ") or t.startswith("Figure ")) and ":" in t[:12]:
        p.paragraph_format.keep_with_next = True
# remove page numbers if any (none in reference) and save
OUT = "Knin_THM_submission_anonymous.docx" if ANON else "Knin_THM_submission.docx"
cp = doc.core_properties; cp.title = TITLE.title(); cp.comments = ""
cp.author = "" if ANON else "Vedran Uroš, Marijo Ćaćić, Ivan Sorić"; cp.last_modified_by = ""
doc.save(OUT); os.remove("_thm_body.docx")
print("written paper/" + OUT)
